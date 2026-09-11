from __future__ import annotations

import hashlib
import html
from contextlib import suppress
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from jinja2 import (
    BaseLoader,
    ChoiceLoader,
    Environment,
    FileSystemLoader,
    PackageLoader,
    select_autoescape,
)
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from pm2.application.artifacts import ArtifactDataService
from pm2.application.services import project_counts
from pm2.application.validation import ValidationService
from pm2.infrastructure.orm import (
    AcceptanceCriterionModel,
    AcceptanceTestModel,
    ChangeModel,
    CommunicationModel,
    DecisionModel,
    DeliverableModel,
    DocumentModel,
    DocumentVersionModel,
    ImplementationActivityModel,
    IssueModel,
    LessonLearnedModel,
    MeetingModel,
    ProjectModel,
    QualityControlModel,
    RequirementModel,
    RiskModel,
    StakeholderModel,
    TransitionActivityModel,
    WbsNodeModel,
)
from pm2.methodology.models import ArtifactDefinition, PM2Configuration


@dataclass(frozen=True, slots=True)
class GeneratedArtifact:
    path: Path
    artifact_code: str
    format: str
    sha256: str


class DocumentService:
    def __init__(
        self, session: Session, methodology: PM2Configuration, template_dir: Path | None = None
    ) -> None:
        self.session = session
        self.methodology = methodology
        if template_dir:
            self.environment = Environment(
                loader=FileSystemLoader(template_dir),
                autoescape=select_autoescape(["html"]),
                trim_blocks=True,
                lstrip_blocks=True,
            )
        else:
            loaders: list[BaseLoader] = [
                FileSystemLoader(Path(__file__).resolve().parents[3] / "templates")
            ]
            with suppress(ValueError):
                loaders.insert(0, PackageLoader("pm2", "templates"))
            self.environment = Environment(
                loader=ChoiceLoader(loaders),
                autoescape=select_autoescape(["html"]),
                trim_blocks=True,
                lstrip_blocks=True,
            )

    def ensure_catalog(self, project_id: str) -> list[DocumentModel]:
        existing = {
            item.artifact_code: item
            for item in self.session.scalars(
                select(DocumentModel).where(
                    DocumentModel.project_id == project_id,
                    DocumentModel.document_type == "PM2_ARTIFACT",
                )
            )
        }
        result: list[DocumentModel] = []
        for code, definition in self.methodology.artifacts.items():
            document = existing.get(code)
            if document is None:
                document = DocumentModel(
                    project_id=project_id,
                    code=code,
                    title=definition.name,
                    document_type="PM2_ARTIFACT",
                    artifact_code=code,
                    status="DRAFT",
                )
                self.session.add(document)
            result.append(document)
        self.session.flush()
        return result

    def context(self, project_id: str, artifact_code: str) -> dict[str, Any]:
        project = self.session.get(ProjectModel, project_id)
        if project is None:
            raise LookupError(f"Projet introuvable : {project_id}")
        try:
            definition: ArtifactDefinition = self.methodology.artifacts[artifact_code]
        except KeyError as exc:
            raise ValueError(f"Artefact PM² inconnu : {artifact_code}") from exc

        def rows(model: Any) -> list[dict[str, Any]]:
            order_column = getattr(model, "code", getattr(model, "name", model.id))
            values = self.session.scalars(
                select(model).where(model.project_id == project_id).order_by(order_column)
            ).all()
            return [
                {column.name: getattr(value, column.name) for column in value.__table__.columns}
                for value in values
            ]

        generated_at = datetime.now(UTC).replace(microsecond=0).isoformat()
        document = self.session.scalar(
            select(DocumentModel).where(
                DocumentModel.project_id == project_id,
                DocumentModel.artifact_code == artifact_code,
            )
        )
        version_count = (
            int(
                self.session.scalar(
                    select(func.count())
                    .select_from(DocumentVersionModel)
                    .where(DocumentVersionModel.document_id == document.id)
                )
                or 0
            )
            if document
            else 0
        )
        artifact_data = ArtifactDataService(self.session).load(project_id, artifact_code)
        schema = ArtifactDataService.schema(artifact_code)
        artifact_sections = [
            {
                "title": section.title,
                "fields": [
                    {
                        "code": field.code,
                        "label": field.label,
                        "value": artifact_data.get(field.code, ""),
                        "required": field.required,
                    }
                    for field in section.fields
                ],
            }
            for section in schema.sections
        ]
        criteria = self.session.scalars(
            select(AcceptanceCriterionModel)
            .join(DeliverableModel)
            .where(DeliverableModel.project_id == project_id)
            .order_by(AcceptanceCriterionModel.code)
        ).all()
        tests = self.session.scalars(
            select(AcceptanceTestModel)
            .join(AcceptanceCriterionModel)
            .join(DeliverableModel)
            .where(DeliverableModel.project_id == project_id)
            .order_by(AcceptanceTestModel.code)
        ).all()

        def to_dict(value: Any) -> dict[str, Any]:
            return {column.name: getattr(value, column.name) for column in value.__table__.columns}

        return {
            "project": {
                "id": project.id,
                "reference": project.reference,
                "name": project.name,
                "description": project.description,
                "status": project.status,
                "phase": project.current_phase,
                "budget": project.approved_budget,
                "currency": project.currency,
                "start_date": project.start_date,
                "target_end_date": project.target_end_date,
                "project_manager": project.project_manager,
                "business_owner": project.business_owner,
            },
            "artifact": {
                "code": artifact_code,
                "title": definition.name,
                "phase": definition.phase,
                "required": definition.required,
                "version": f"1.{version_count}",
                "status": document.status if document else "DRAFT",
                "author": document.author if document else project.project_manager,
                "reviewer": document.reviewer if document else None,
                "approver": document.approver if document else None,
                "approval_date": document.approval_date if document else None,
                "created_at": document.created_at if document else None,
                "updated_at": document.updated_at if document else None,
                "baseline": bool(document and document.status == "BASELINED"),
            },
            "generated_at": generated_at,
            "artifact_data": artifact_data,
            "artifact_sections": artifact_sections,
            "artifact_purpose": schema.purpose,
            "counts": project_counts(self.session, project_id),
            "stakeholders": rows(StakeholderModel),
            "wbs": rows(WbsNodeModel),
            "requirements": rows(RequirementModel),
            "deliverables": rows(DeliverableModel),
            "risks": rows(RiskModel),
            "issues": rows(IssueModel),
            "decisions": rows(DecisionModel),
            "changes": rows(ChangeModel),
            "meetings": rows(MeetingModel),
            "quality_controls": rows(QualityControlModel),
            "transition_activities": rows(TransitionActivityModel),
            "implementation_activities": rows(ImplementationActivityModel),
            "communications": rows(CommunicationModel),
            "lessons_learned": rows(LessonLearnedModel),
            "acceptance_criteria": [to_dict(value) for value in criteria],
            "acceptance_tests": [to_dict(value) for value in tests],
            "validations": [
                problem.__dict__
                if hasattr(problem, "__dict__")
                else {
                    "code": problem.code,
                    "severity": problem.severity,
                    "message": problem.message,
                    "entity": problem.entity,
                    "entity_id": problem.entity_id,
                }
                for problem in ValidationService(self.session).validate_project(project_id)
            ],
        }

    def render_markdown(self, context: dict[str, Any]) -> str:
        name = f"artifacts/{context['artifact']['code']}.md.j2"
        return self.environment.get_template(name).render(**context).strip() + "\n"

    def render_html(self, context: dict[str, Any]) -> str:
        name = f"artifacts/{context['artifact']['code']}.html.j2"
        return self.environment.overlay(autoescape=True).get_template(name).render(**context)

    def generate(
        self, project_id: str, artifact_code: str, output: Path, format: str
    ) -> GeneratedArtifact:
        fmt = format.lower().lstrip(".")
        if fmt not in {"md", "markdown", "html", "docx", "pdf"}:
            raise ValueError(f"Format d'export non pris en charge : {format}")
        output.parent.mkdir(parents=True, exist_ok=True)
        context = self.context(project_id, artifact_code)
        if fmt in {"md", "markdown"}:
            output.write_text(self.render_markdown(context), encoding="utf-8")
            fmt = "md"
        elif fmt == "html":
            output.write_text(self.render_html(context), encoding="utf-8")
        elif fmt == "docx":
            self._write_docx(output, context)
        else:
            self._write_pdf(output, context)
        digest = hashlib.sha256(output.read_bytes()).hexdigest()
        self._record_version(project_id, artifact_code, output, digest)
        return GeneratedArtifact(output, artifact_code, fmt, digest)

    def _record_version(
        self, project_id: str, artifact_code: str, output: Path, digest: str
    ) -> None:
        document = self.session.scalar(
            select(DocumentModel).where(
                DocumentModel.project_id == project_id,
                DocumentModel.artifact_code == artifact_code,
            )
        )
        if document is None:
            self.ensure_catalog(project_id)
            document = self.session.scalar(
                select(DocumentModel).where(
                    DocumentModel.project_id == project_id,
                    DocumentModel.artifact_code == artifact_code,
                )
            )
        assert document is not None
        version_count = len(
            self.session.scalars(
                select(DocumentVersionModel).where(DocumentVersionModel.document_id == document.id)
            ).all()
        )
        version = DocumentVersionModel(
            document_id=document.id,
            version=f"1.{version_count}",
            status=document.status,
            content_hash=digest,
            file_path=str(output),
            baseline=document.status == "BASELINED",
        )
        self.session.add(version)
        self.session.flush()

    @staticmethod
    def _write_docx(path: Path, context: dict[str, Any]) -> None:
        document = DocxDocument()
        styles = document.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(10)
        title = document.add_heading(context["artifact"]["title"], 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph(
            f"{context['project']['reference']} — {context['project']['name']}"
        ).alignment = WD_ALIGN_PARAGRAPH.CENTER
        metadata = document.add_table(rows=0, cols=2)
        metadata.style = "Table Grid"
        for label, value in (
            ("Code", context["artifact"]["code"]),
            ("Version", context["artifact"]["version"]),
            ("Statut", context["artifact"]["status"]),
            ("Auteur", context["artifact"]["author"] or "—"),
            ("Réviseur", context["artifact"]["reviewer"] or "—"),
            ("Approbateur", context["artifact"]["approver"] or "—"),
            ("Créé le", context["artifact"]["created_at"] or "—"),
            ("Modifié le", context["artifact"]["updated_at"] or "—"),
            ("Approuvé le", context["artifact"]["approval_date"] or "—"),
            ("Référence de base", "Oui" if context["artifact"]["baseline"] else "Non"),
            ("Généré le", context["generated_at"]),
        ):
            cells = metadata.add_row().cells
            cells[0].text, cells[1].text = label, str(value)
        document.add_heading("Projet", level=1)
        document.add_paragraph(context["project"]["description"] or "Aucune description.")
        for section in context["artifact_sections"]:
            document.add_heading(section["title"], level=1)
            for item in section["fields"]:
                document.add_heading(item["label"], level=2)
                document.add_paragraph(item["value"] or "Non renseigné.")
        DocumentService._append_docx_data(document, context)
        document.add_heading("Validation PM²", level=1)
        if context["validations"]:
            for problem in context["validations"]:
                document.add_paragraph(
                    f"[{problem['severity']}] {problem['code']} — {problem['message']}",
                    style="List Bullet",
                )
        else:
            document.add_paragraph("Aucun problème détecté.")
        document.save(str(path))

    @staticmethod
    def _append_docx_data(document: DocxDocumentObject, context: dict[str, Any]) -> None:
        for key, title in (
            ("requirements", "Exigences"),
            ("deliverables", "Livrables"),
            ("risks", "Risques"),
            ("issues", "Problèmes"),
            ("decisions", "Décisions"),
            ("changes", "Modifications"),
        ):
            document.add_heading(title, level=1)
            items = context[key]
            if not items:
                document.add_paragraph("Aucun élément.")
                continue
            table = document.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            table.rows[0].cells[0].text = "Code"
            table.rows[0].cells[1].text = "Titre"
            table.rows[0].cells[2].text = "Statut"
            for item in items:
                cells = table.add_row().cells
                cells[0].text = str(item.get("code", ""))
                cells[1].text = str(item.get("title", item.get("name", "")))
                cells[2].text = str(item.get("status", ""))

    @staticmethod
    def _write_pdf(path: Path, context: dict[str, Any]) -> None:
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "PM2Title", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=12
        )
        story: list[Any] = [
            Paragraph(html.escape(str(context["artifact"]["title"])), title_style),
            Paragraph(
                html.escape(f"{context['project']['reference']} — {context['project']['name']}"),
                styles["Heading2"],
            ),
            Spacer(1, 5 * mm),
        ]
        metadata = [
            ["Code", context["artifact"]["code"]],
            ["Version", context["artifact"]["version"]],
            ["Statut", context["artifact"]["status"]],
            ["Auteur", context["artifact"]["author"] or "—"],
            ["Réviseur", context["artifact"]["reviewer"] or "—"],
            ["Approbateur", context["artifact"]["approver"] or "—"],
            ["Créé le", context["artifact"]["created_at"] or "—"],
            ["Modifié le", context["artifact"]["updated_at"] or "—"],
            ["Approuvé le", context["artifact"]["approval_date"] or "—"],
            ["Référence de base", "Oui" if context["artifact"]["baseline"] else "Non"],
            ["Généré le", context["generated_at"]],
        ]
        table = Table(metadata, colWidths=[35 * mm, 120 * mm])
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E8EEF8")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.extend(
            [
                table,
                Spacer(1, 7 * mm),
                Paragraph("Projet", styles["Heading1"]),
                Paragraph(
                    html.escape(context["project"]["description"] or "Aucune description."),
                    styles["BodyText"],
                ),
            ]
        )
        for section in context["artifact_sections"]:
            story.append(Paragraph(html.escape(section["title"]), styles["Heading1"]))
            for item in section["fields"]:
                story.append(Paragraph(html.escape(item["label"]), styles["Heading2"]))
                story.append(
                    Paragraph(html.escape(item["value"] or "Non renseigné."), styles["BodyText"])
                )
        for key, title in (
            ("requirements", "Exigences"),
            ("deliverables", "Livrables"),
            ("risks", "Risques"),
            ("issues", "Problèmes"),
            ("decisions", "Décisions"),
            ("changes", "Modifications"),
        ):
            story.append(Paragraph(title, styles["Heading1"]))
            data = [["Code", "Titre", "Statut"]]
            for item in context[key]:
                data.append(
                    [
                        str(item.get("code", "")),
                        Paragraph(
                            html.escape(str(item.get("title", item.get("name", "")))),
                            styles["BodyText"],
                        ),
                        str(item.get("status", "")),
                    ]
                )
            if len(data) == 1:
                story.append(Paragraph("Aucun élément.", styles["BodyText"]))
            else:
                item_table = Table(data, colWidths=[30 * mm, 100 * mm, 35 * mm], repeatRows=1)
                item_table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#254E70")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                story.append(item_table)
        story.extend([PageBreak(), Paragraph("Validation PM²", styles["Heading1"])])
        if context["validations"]:
            for problem in context["validations"]:
                story.append(
                    Paragraph(
                        html.escape(
                            f"[{problem['severity']}] {problem['code']} — {problem['message']}"
                        ),
                        styles["BodyText"],
                    )
                )
        else:
            story.append(Paragraph("Aucun problème détecté.", styles["BodyText"]))
        pdf = SimpleDocTemplate(
            str(path),
            pagesize=A4,
            rightMargin=15 * mm,
            leftMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
            title=context["artifact"]["title"],
            author="PM² Desktop",
        )
        pdf.build(story)
